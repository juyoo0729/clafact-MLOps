import csv
import json

from openpyxl import Workbook, load_workbook

from tools.run_r3_registry_full_value_replay import run


def _write_ledger(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1542건 원장"
    sheet.append([
        "claim_id", "article_id", "article_date", "split", "claim_type_original",
        "gold_time_source", "sentence", "indicator", "value", "unit", "time",
        "frequency", "region", "population", "dimension", "comparison",
        "calculation", "condition", "source_hint", "original_parse_status",
        "missing_required_slots", "flexible_route", "route_reason", "next_action",
        "target_value_role_status", "runtime_r3_admission",
    ])
    for claim_id in ("A1_1", "A1_2"):
        sheet.append([
            claim_id, "A1", "2025-01-01", "dev", "DIRECT_VALUE", "explicit",
            "SENTINEL_ARTICLE_TEXT_MUST_NOT_APPEAR", "취업자 수", "28000", "천명",
            "2025년 1월", "월", "전국", None, "{}", "{}", "DIRECT_VALUE",
            "{}", None, "AUTO_OK", "", "R2_SLOT_READY", "", "", "MISSING",
            "READY",
        ])
    workbook.save(path)


def _write_concepts(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "03_1542_Concept_전체"
    sheet.append([
        "article_id", "sentence_id", "sentence", "indicator", "value", "unit",
        "time", "frequency", "region", "population", "dimension", "comparison",
        "calculation", "condition", "source_hint", "parse_status", "concept_id",
        "canonical_name", "standard_key", "concept_action", "concept_status",
        "concept_reason", "kosis_fit_status",
    ])
    for sentence_id in ("1", "2"):
        sheet.append([
            "A1", sentence_id, "SENTINEL_CONCEPT_SENTENCE_MUST_NOT_APPEAR",
            "취업자 수", "28000", "천명", "2025년 1월", "월", "전국", None,
            "{}", "{}", "DIRECT_VALUE", "{}", None, "AUTO_OK", "C1",
            "취업자", "EMPLOYED_PERSON_COUNT", "EXISTING", "ASSIGNED", "",
            "KOSIS_CANDIDATE",
        ])
    workbook.save(path)


def test_full_runner_joins_all_claims_and_deduplicates_value_calls(tmp_path):
    ledger = tmp_path / "ledger.xlsx"
    concepts = tmp_path / "concepts.xlsx"
    candidates = tmp_path / "candidates.csv"
    identities = tmp_path / "identities.jsonl"
    catalog = tmp_path / "catalog.json"
    coordinates = tmp_path / "coordinates.json"
    members = tmp_path / "members.json"
    output = tmp_path / "output"
    _write_ledger(ledger)
    _write_concepts(concepts)
    with candidates.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["claim_id", "candidate_tbl_ids"])
        writer.writeheader()
        writer.writerows([
            {"claim_id": "A1_1", "candidate_tbl_ids": "DT_EMP"},
            {"claim_id": "A1_2", "candidate_tbl_ids": "DT_EMP"},
        ])
    identities.write_text(
        json.dumps({
            "indicator": "취업자 수",
            "candidates": [
                {"org_id": "101", "tbl_id": "DT_EMP", "tbl_name": "취업자"}
            ],
        }, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    catalog.write_text("[]\n", encoding="utf-8")
    coordinates.write_text(json.dumps([{
        "tbl_id": "DT_EMP", "itm_id": "T30",
        "indicator_aliases": ["취업자 수"],
        "dimension_members": {"B": "계", "J": "계"},
    }], ensure_ascii=False), encoding="utf-8")
    members.write_text(json.dumps([
        {"tbl_id": "DT_EMP", "dimension_id": "B", "member_name": "계", "member_code": "0"},
        {"tbl_id": "DT_EMP", "dimension_id": "J", "member_name": "계", "member_code": "00"},
    ], ensure_ascii=False), encoding="utf-8")

    value_calls = []

    def metadata_fetcher(api_key, org_id, table_id, *, meta_type):
        if meta_type == "ITM":
            return [
                {"ORG_ID": org_id, "TBL_ID": table_id, "OBJ_ID": "item", "ITM_ID": "T30", "ITM_NM": "취업자", "UNIT_NM": "천명", "OBJ_ID_SN": "0"},
                {"ORG_ID": org_id, "TBL_ID": table_id, "OBJ_ID": "B", "OBJ_NM": "성별", "ITM_ID": "0", "ITM_NM": "계", "OBJ_ID_SN": "1"},
                {"ORG_ID": org_id, "TBL_ID": table_id, "OBJ_ID": "J", "OBJ_NM": "종사상지위별", "ITM_ID": "00", "ITM_NM": "계", "OBJ_ID_SN": "2"},
            ]
        return [{"PRD_SE": "월", "STRT_PRD_DE": "2025.01", "END_PRD_DE": "2025.12"}]

    def value_fetcher(api_key, org_id, table_id, item_id, period_type, start_period, end_period, object_codes):
        value_calls.append((org_id, table_id, item_id, period_type, start_period, tuple(object_codes)))
        return [{
            "ORG_ID": org_id, "TBL_ID": table_id, "ITM_ID": item_id,
            "PRD_SE": period_type, "PRD_DE": start_period,
            "C1": object_codes[0], "C2": object_codes[1],
            "DT": "28000", "UNIT_NM": "천명",
        }]

    summary = run(
        ledger_xlsx=ledger,
        concept_xlsx=concepts,
        candidate_attachment_csv=candidates,
        candidate_identity_jsonl=identities,
        catalog_json=catalog,
        registered_coordinates_json=coordinates,
        member_codes_json=members,
        output_dir=output,
        api_key="SECRET_MUST_NOT_BE_RECORDED",
        metadata_fetcher=metadata_fetcher,
        value_fetcher=value_fetcher,
        requests_per_minute=0,
    )

    assert summary["input_claim_count"] == 2
    assert summary["official_value_linked_claim_count"] == 2
    assert summary["unique_value_api_call_count"] == 1
    assert len(value_calls) == 1
    with (output / "claim_value_results.csv").open(encoding="utf-8-sig") as handle:
        assert len(list(csv.DictReader(handle))) == 2
    workbook = load_workbook(output / "CLAFACT_1542_공식값_조회결과.xlsx", read_only=True)
    assert workbook.sheetnames == ["요약", "Registry", "Claim 공식값", "HOLD 사유"]
    combined = "\n".join(
        path.read_text(encoding="utf-8-sig")
        for path in output.iterdir()
        if path.suffix in {".csv", ".json", ".jsonl"}
    )
    assert "SENTINEL_ARTICLE_TEXT_MUST_NOT_APPEAR" not in combined
    assert "SENTINEL_CONCEPT_SENTENCE_MUST_NOT_APPEAR" not in combined
    assert "SECRET_MUST_NOT_BE_RECORDED" not in combined

    replay_output = tmp_path / "replay_output"

    def unexpected_metadata_call(*args, **kwargs):
        raise AssertionError("cached official metadata must not be called again")

    def unexpected_value_call(*args, **kwargs):
        raise AssertionError("cached official cell must not be called again")

    replay_summary = run(
        ledger_xlsx=ledger,
        concept_xlsx=concepts,
        candidate_attachment_csv=candidates,
        candidate_identity_jsonl=identities,
        catalog_json=catalog,
        registered_coordinates_json=coordinates,
        member_codes_json=members,
        output_dir=replay_output,
        api_key="",
        allow_live_kosis=False,
        metadata_fetcher=unexpected_metadata_call,
        value_fetcher=unexpected_value_call,
        metadata_cache_paths=[output / "metadata_snapshots.jsonl"],
        value_cache_paths=[output / "value_snapshots.jsonl"],
        requests_per_minute=0,
    )

    assert replay_summary["official_value_linked_claim_count"] == 2
    assert replay_summary["metadata_live_api_call_count"] == 0
    assert replay_summary["unique_value_api_call_count"] == 0
    assert replay_summary["value_cache_hit_count"] == 1
    manifest = json.loads((replay_output / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["secrets"]["kosis_api_key"] == "ABSENT_OFFLINE_CACHE_ONLY"

    missing_value_cache = tmp_path / "missing_value_cache.jsonl"
    missing_value_cache.write_text("", encoding="utf-8")
    missing_output = tmp_path / "missing_output"
    missing_summary = run(
        ledger_xlsx=ledger,
        concept_xlsx=concepts,
        candidate_attachment_csv=candidates,
        candidate_identity_jsonl=identities,
        catalog_json=catalog,
        registered_coordinates_json=coordinates,
        member_codes_json=members,
        output_dir=missing_output,
        api_key="",
        allow_live_kosis=False,
        metadata_fetcher=unexpected_metadata_call,
        value_fetcher=unexpected_value_call,
        metadata_cache_paths=[output / "metadata_snapshots.jsonl"],
        value_cache_paths=[missing_value_cache],
        requests_per_minute=0,
    )

    assert missing_summary["unique_value_api_call_count"] == 0
    assert missing_summary["value_offline_cache_miss_count"] == 1
    assert missing_summary["official_value_linked_claim_count"] == 0
