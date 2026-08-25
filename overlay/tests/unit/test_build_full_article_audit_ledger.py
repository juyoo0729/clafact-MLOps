import csv
import json

from openpyxl import load_workbook

from tools.build_full_article_audit_ledger import build, refresh_output_manifest


def _write_csv(path, fieldnames, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def test_build_joins_full_article_claim_and_kosis_trace(tmp_path):
    articles = tmp_path / "articles.csv"
    ledger = tmp_path / "ledger.csv"
    subtypes = tmp_path / "subtypes.csv"
    values = tmp_path / "values.csv"
    metadata = tmp_path / "metadata.jsonl"
    snapshots = tmp_path / "values.jsonl"
    replay_summary = tmp_path / "summary.json"

    sentence = "지난달 취업자 수는 10만 명 증가했다."
    body = f"앞 문장이다. {sentence} 뒤 문장이다."
    _write_csv(
        articles,
        ["번호", "작성일", "제목", "URL", "검색구분레이블", "원본CSV행", "파일명", "본문"],
        [{
            "번호": "1", "작성일": "2025-02-14", "제목": "고용 기사",
            "URL": "https://example.test/article", "검색구분레이블": "True",
            "원본CSV행": "2", "파일명": "0001.txt", "본문": body,
        }],
    )
    _write_csv(
        ledger,
        [
            "기사번호", "문장번호", "부모Claim번호", "Claim번호", "원문",
            "12개항목상태", "12개항목공식조회가능", "통계표검색시도",
            "항목정보조회시도", "기간정보조회시도", "현재상태",
            "현재중단단계", "현재사유", "대표문제", "보조문제",
            "다음실행단계", "최신결과상태", "최신결과단계",
            "최신결과사유", "최신개선판정",
        ],
        [{
            "기사번호": "A00001", "문장번호": "2", "부모Claim번호": "A00001_2",
            "Claim번호": "A00001_2", "원문": sentence, "현재상태": "AUTO",
        }],
    )
    _write_csv(
        subtypes,
        ["탭", "Claim번호", "원문", "하위유형", "속성", "분류근거"],
        [{
            "탭": "08_직접값", "Claim번호": "A00001_2", "원문": sentence,
            "하위유형": "직접값", "속성": "{}", "분류근거": "test",
        }],
    )
    _write_csv(
        values,
        [
            "claim_id", "concept_id", "standard_key", "indicator", "claim_unit",
            "claim_time", "claim_frequency", "calculation_type", "registry_signature",
            "candidate_count", "coordinate_status", "reason_code",
            "coordinate_provenance", "org_id", "table_id", "item_id",
            "object_codes", "period_type", "period", "official_value_status",
            "official_value", "official_unit", "verdict_status",
        ],
        [{
            "claim_id": "A00001_2", "concept_id": "C1", "standard_key": "EMPLOYMENT",
            "indicator": "취업자 수", "claim_unit": "명", "claim_time": "2025-01",
            "claim_frequency": "월", "calculation_type": "DIRECT_VALUE",
            "registry_signature": "employment", "candidate_count": "1",
            "coordinate_status": "REGISTERED_COORDINATE_READY", "reason_code": "",
            "coordinate_provenance": "REGISTERED_COORDINATE", "org_id": "101",
            "table_id": "DT_TEST", "item_id": "T1", "object_codes": "0 | 00",
            "period_type": "M", "period": "202501",
            "official_value_status": "OFFICIAL_VALUE_FETCHED", "official_value": "10",
            "official_unit": "명", "verdict_status": "NOT_EVALUATED",
        }],
    )
    metadata.write_text(
        json.dumps({
            "request": {"org_id": "101", "table_id": "DT_TEST", "meta_type": "ITM"},
            "retrieved_at": "2025-02-15T00:00:00+00:00", "retrieval_source": "LIVE_KOSIS",
            "status": "SUCCESS", "error_code": "", "response_sha256": "meta-hash",
            "response": [{"OBJ_NM": "항목", "ITM_NM": "취업자", "ITM_ID": "T1"}],
        }, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    snapshots.write_text(
        json.dumps({
            "request": {
                "org_id": "101", "table_id": "DT_TEST", "item_id": "T1",
                "period_type": "M", "start_period": "202501", "end_period": "202501",
                "object_codes": ["0", "00"],
            },
            "retrieved_at": "2025-02-15T00:00:01+00:00", "retrieval_source": "LIVE_KOSIS",
            "status": "SUCCESS", "error_code": "", "response_sha256": "value-hash",
            "response": [{
                "TBL_NM": "고용", "ITM_NM": "취업자", "PRD_DE": "202501",
                "DT": "10", "UNIT_NM": "명",
            }],
        }, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    replay_summary.write_text(
        json.dumps({"verdict_status": "NOT_EVALUATED", "accuracy_status": "NOT_EVALUABLE"}),
        encoding="utf-8",
    )

    output = tmp_path / "output"
    summary = build(
        articles_csv=articles,
        claim_ledger_csv=ledger,
        subtype_csv=subtypes,
        claim_values_csv=values,
        metadata_snapshots_jsonl=metadata,
        value_snapshots_jsonl=snapshots,
        replay_summary_json=replay_summary,
        output_dir=output,
        expected_article_count=1,
        expected_claim_count=1,
    )

    assert summary["article_count"] == 1
    assert summary["claim_exact_context_join_count"] == 1
    assert summary["official_value_linked_claim_count"] == 1
    assert summary["metadata_query_record_count"] == 1
    assert summary["value_query_record_count"] == 1

    workbook = load_workbook(output / "CLAFACT_499기사_1542Claim_전수실행감사원장.xlsx")
    assert workbook.sheetnames == [
        "요약", "전체 기사", "Claim 실행원장", "KOSIS 조회기록", "사유 집계", "실행정보",
    ]
    assert workbook["Claim 실행원장"].max_row == 2
    assert workbook["요약"]["B3"].data_type == "f"

    manifest = refresh_output_manifest(output)
    assert manifest["outputs"]["claim_audit_1542.csv"]["bytes"] > 0
