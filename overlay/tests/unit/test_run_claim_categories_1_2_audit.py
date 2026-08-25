import csv
import json
from datetime import date

from openpyxl import load_workbook

from tools.run_claim_categories_1_2_audit import run, refresh_output_manifest


def _write_csv(path, rows):
    fields = sorted({key for row in rows for key in row})
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _split_runner(method, sentence):
    assert method == "rule_ensemble"
    assert sentence == "취업자는 100만 명이고 실업자는 20만 명이다."
    return {
        "route_status": "AUTO",
        "reason_code": "RULE_SPLIT",
        "children": [
            {"text": "취업자는 100만 명이다.", "target_value_text": "100만 명"},
            {"text": "실업자는 20만 명이다.", "target_value_text": "20만 명"},
        ],
    }


def test_run_records_category_one_and_two_with_six_w_audit(tmp_path):
    claim_audit = tmp_path / "claim_audit.csv"
    instruction = tmp_path / "instructions.txt"
    output = tmp_path / "output"
    _write_csv(
        claim_audit,
        [
            {
                "분류탭": "01_문맥보완",
                "기사번호": "A00001",
                "Claim번호": "A00001_1",
                "작성일": "2025-02-14",
                "제목": "고용 기사",
                "URL": "https://example.test/1",
                "원문": "지난달 취업자는 10만 명 증가했다.",
                "앞문맥": "통계청이 고용동향을 발표했다.",
                "뒤문맥": "증가세가 이어졌다.",
                "하위유형": "시점문맥필요형",
                "기사원문정확일치": "YES",
                "기사내원문시작위치": "16",
                "공식값상태": "HOLD",
                "최종성공실패사유": "PERIOD_MISSING",
            },
            {
                "분류탭": "02_복수Claim분리",
                "기사번호": "A00002",
                "Claim번호": "A00002_1",
                "작성일": "2025-02-14",
                "제목": "고용 기사 2",
                "URL": "https://example.test/2",
                "원문": "취업자는 100만 명이고 실업자는 20만 명이다.",
                "앞문맥": "통계청 발표다.",
                "뒤문맥": "두 지표를 함께 설명했다.",
                "하위유형": "병렬수치형",
                "기사원문정확일치": "YES",
                "기사내원문시작위치": "20",
                "공식값상태": "HOLD",
                "최종성공실패사유": "MULTI_VALUE",
            },
        ],
    )
    instruction.write_text(
        "1. 문맥 보완 필요 — 1건\n2. 복수 Claim 분리 필요 — 1건\n",
        encoding="utf-8",
    )

    summary = run(
        claim_audit_csv=claim_audit,
        instruction_txt=instruction,
        output_dir=output,
        split_runner=_split_runner,
        expected_category1_count=1,
        expected_category2_count=1,
    )

    assert summary["actual_category1_count"] == 1
    assert summary["context_status_counts"] == {"SUCCESS": 1}
    assert summary["split_parent_status_counts"] == {"SUCCESS": 1}
    assert summary["split_child_count"] == 2
    assert summary["split_child_status_counts"] == {"VALID": 2}
    assert summary["split_safe_child_count"] == 2
    assert summary["event_count"] == 4
    assert summary["kosis_requery_count"] == 0

    context_rows = list(csv.DictReader(
        (output / "category1_context_results.csv").open(encoding="utf-8-sig")
    ))
    assert context_rows[0]["보완기준기간"] == "2025-01"
    assert context_rows[0]["누가"]
    assert context_rows[0]["언제"]
    assert context_rows[0]["어디서"]
    assert context_rows[0]["무엇을"]
    assert context_rows[0]["어떻게"]
    assert context_rows[0]["왜"]

    events = [
        json.loads(line)
        for line in (output / "execution_events.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    assert len(events) == 4
    assert {event["이벤트종류"] for event in events} == {
        "CONTEXT_PARENT", "SPLIT_PARENT", "SPLIT_CHILD",
    }
    text_report = (output / "CLAFACT_1번2번_육하원칙_실행기록.txt").read_text(
        encoding="utf-8"
    )
    assert "육하원칙 상세 이벤트: 4건" in text_report
    assert "누가: 기사 A00001 / Claim A00001_1" in text_report
    assert "어떻게: ARTICLE_BODY_EXACT_JOIN" in text_report

    consolidated = json.loads(
        (output / "CLAFACT_1번2번_595건_육하원칙_통합기록.json").read_text(
            encoding="utf-8"
        )
    )
    assert consolidated["scope"]["actual_parent_claim_count"] == 2
    assert len(consolidated["category1_context_results"]) == 1
    assert len(consolidated["category2"]["parent_results"]) == 1
    assert len(consolidated["category2"]["child_results"]) == 2
    assert len(consolidated["six_w_execution_events"]) == 4

    workbook = load_workbook(output / "CLAFACT_1번2번_595건_육하원칙_실행원장.xlsx")
    assert workbook.sheetnames == [
        "요약", "1_문맥보완", "2_부모Claim", "2_자식Claim",
        "육하원칙 이벤트", "성공실패 사유", "실행정보",
    ]
    assert workbook["요약"]["B3"].value == 1
    assert workbook["요약"]["B4"].data_type == "f"
    assert workbook["요약"]["B15"].data_type == "f"

    manifest = refresh_output_manifest(output)
    assert manifest["outputs"]["execution_events.jsonl"]["bytes"] > 0


def test_context_period_prefers_relative_target_and_rejects_conflicting_context():
    from tools.run_claim_categories_1_2_audit import resolve_period

    resolved = resolve_period(
        "내수 판매가 감소했다.",
        "작년 판매량이 2023년 대비 줄었다.",
        "작년 판매량은 100만 대였다.",
        date(2025, 1, 3),
    )
    assert resolved["target_period"] == "2024"
    assert resolved["comparison_period"] == "2023"
    assert resolved["reason_code"] == "CONSISTENT_PERIOD_RESOLVED_FROM_BOTH_CONTEXT_SIDES"

    conflict = resolve_period(
        "판매가 감소했다.",
        "2023년 판매 실적이다.",
        "2024년 판매 실적이다.",
        date(2025, 1, 3),
    )
    assert conflict["target_period"] == ""
    assert conflict["reason_code"] == "MULTIPLE_CONTEXT_PERIODS_UNRESOLVED"

    absolute_candidate = resolve_period(
        "판매가 감소했다.",
        "2023년 판매 실적이다.",
        "",
        date(2025, 1, 3),
    )
    assert absolute_candidate["target_period"] == "2023"
    assert absolute_candidate["reason_code"] == (
        "CONTEXT_ABSOLUTE_PERIOD_CANDIDATE_REQUIRES_SEMANTIC_LINK"
    )
