"""Build one immutable audit ledger for full articles, Claims, and KOSIS calls."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ARTICLE_COLUMNS = (
    "기사번호",
    "작성일",
    "제목",
    "URL",
    "검색구분레이블",
    "원본CSV행",
    "파일명",
    "본문길이",
    "본문SHA256",
    "Claim수",
    "Claim원문정확일치수",
    "기사문맥연결상태",
    "원문데이터파일",
    "본문미리보기",
)

FULL_ARTICLE_COLUMNS = ARTICLE_COLUMNS[:-1] + ("본문",)

CLAIM_COLUMNS = (
    "기사번호",
    "Claim번호",
    "문장번호",
    "부모Claim번호",
    "작성일",
    "제목",
    "URL",
    "원문",
    "기사내원문시작위치",
    "기사원문정확일치",
    "앞문맥",
    "뒤문맥",
    "분류탭",
    "하위유형",
    "분류속성",
    "분류근거",
    "12개항목상태",
    "12개항목공식조회가능",
    "통계표검색시도",
    "항목정보조회시도",
    "기간정보조회시도",
    "현재상태",
    "현재중단단계",
    "현재사유",
    "대표문제",
    "보조문제",
    "다음실행단계",
    "최신결과상태",
    "최신결과단계",
    "최신결과사유",
    "최신개선판정",
    "concept_id",
    "standard_key",
    "지표",
    "Claim단위",
    "Claim기간",
    "Claim주기",
    "계산유형",
    "registry_signature",
    "후보표수",
    "좌표상태",
    "좌표사유",
    "좌표출처",
    "기관ID",
    "통계표ID",
    "항목ID",
    "분류코드",
    "기간유형",
    "조회기간",
    "KOSIS값조회시도",
    "KOSIS조회시각",
    "KOSIS조회경로",
    "KOSIS조회상태",
    "KOSIS오류코드",
    "KOSIS응답SHA256",
    "공식값상태",
    "공식값",
    "공식단위",
    "Verdict상태",
    "최종실행상태",
    "최종성공실패사유",
    "최종중단단계",
)

QUERY_COLUMNS = (
    "조회종류",
    "조회키",
    "기관ID",
    "통계표ID",
    "메타유형",
    "항목ID",
    "기간유형",
    "시작기간",
    "종료기간",
    "분류코드",
    "조회시각",
    "조회경로",
    "조회상태",
    "오류코드",
    "응답행수",
    "응답SHA256",
    "연결Claim수",
    "연결Claim번호",
    "응답요약",
)

OUTPUT_NAMES = (
    "CLAFACT_499기사_1542Claim_전수실행감사원장.xlsx",
    "articles_499_full.csv",
    "claim_audit_1542.csv",
    "kosis_query_audit.csv",
    "summary.json",
)


def build(
    *,
    articles_csv: Path,
    claim_ledger_csv: Path,
    subtype_csv: Path,
    claim_values_csv: Path,
    metadata_snapshots_jsonl: Path,
    value_snapshots_jsonl: Path,
    replay_summary_json: Path,
    output_dir: Path,
    expected_article_count: int | None = 499,
    expected_claim_count: int | None = 1542,
) -> dict[str, Any]:
    """Join every Claim to its full article and all stored KOSIS call evidence."""
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(f"output directory is not empty: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)

    article_source_rows = _read_csv(articles_csv)
    ledger_rows = _read_csv(claim_ledger_csv)
    subtype_rows = _read_csv(subtype_csv)
    value_rows = _read_csv(claim_values_csv)
    metadata_rows = _read_jsonl(metadata_snapshots_jsonl)
    value_snapshot_rows = _read_jsonl(value_snapshots_jsonl)
    replay_summary = json.loads(replay_summary_json.read_text(encoding="utf-8"))

    _require_unique(ledger_rows, "Claim번호", "Claim ledger")
    _require_unique(subtype_rows, "Claim번호", "subtype ledger")
    _require_unique(value_rows, "claim_id", "official-value ledger")

    claim_ids = {_text(row.get("Claim번호")) for row in ledger_rows}
    if claim_ids != {_text(row.get("Claim번호")) for row in subtype_rows}:
        raise ValueError("subtype Claim IDs do not match the Claim ledger")
    if claim_ids != {_text(row.get("claim_id")) for row in value_rows}:
        raise ValueError("official-value Claim IDs do not match the Claim ledger")

    article_by_id = {
        _article_id(row.get("번호")): row
        for row in article_source_rows
        if _article_id(row.get("번호"))
    }
    article_ids = {_text(row.get("기사번호")) for row in ledger_rows}
    missing_articles = sorted(article_ids - set(article_by_id))
    if missing_articles:
        raise ValueError(
            f"full article bodies missing for {len(missing_articles)} IDs: "
            f"{missing_articles[:10]}"
        )

    if expected_article_count is not None and len(article_ids) != expected_article_count:
        raise ValueError(
            f"unexpected article count: {len(article_ids)} != {expected_article_count}"
        )
    if expected_claim_count is not None and len(ledger_rows) != expected_claim_count:
        raise ValueError(
            f"unexpected Claim count: {len(ledger_rows)} != {expected_claim_count}"
        )

    subtype_by_claim = {
        _text(row.get("Claim번호")): row for row in subtype_rows
    }
    value_by_claim = {_text(row.get("claim_id")): row for row in value_rows}
    value_snapshot_by_key = {
        _snapshot_value_key(row): row for row in value_snapshot_rows
    }

    claims_by_article: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in ledger_rows:
        claims_by_article[_text(row.get("기사번호"))].append(row)

    article_rows = _build_article_rows(
        article_ids=article_ids,
        article_by_id=article_by_id,
        claims_by_article=claims_by_article,
        source_path=articles_csv,
    )
    claim_rows = _build_claim_rows(
        ledger_rows=ledger_rows,
        article_by_id=article_by_id,
        subtype_by_claim=subtype_by_claim,
        value_by_claim=value_by_claim,
        value_snapshot_by_key=value_snapshot_by_key,
    )
    query_rows = _build_query_rows(
        metadata_rows=metadata_rows,
        value_snapshot_rows=value_snapshot_rows,
        claim_rows=claim_rows,
    )

    article_path = output_dir / "articles_499_full.csv"
    claim_path = output_dir / "claim_audit_1542.csv"
    query_path = output_dir / "kosis_query_audit.csv"
    _write_csv(article_path, article_rows, FULL_ARTICLE_COLUMNS)
    _write_csv(claim_path, claim_rows, CLAIM_COLUMNS)
    _write_csv(query_path, query_rows, QUERY_COLUMNS)

    final_reasons = Counter(row["최종성공실패사유"] for row in claim_rows)
    query_statuses = Counter(row["조회상태"] for row in query_rows)
    summary: dict[str, Any] = {
        "artifact": "clafact_full_article_execution_audit_v1",
        "article_count": len(article_rows),
        "claim_count": len(claim_rows),
        "article_body_joined_count": sum(
            row["기사문맥연결상태"] == "COMPLETE" for row in article_rows
        ),
        "claim_exact_context_join_count": sum(
            row["기사원문정확일치"] == "YES" for row in claim_rows
        ),
        "metadata_query_record_count": sum(
            row["조회종류"] == "METADATA" for row in query_rows
        ),
        "value_query_record_count": sum(
            row["조회종류"] == "VALUE" for row in query_rows
        ),
        "query_status_counts": dict(sorted(query_statuses.items())),
        "official_value_linked_claim_count": sum(
            row["최종실행상태"] == "OFFICIAL_VALUE_LINKED" for row in claim_rows
        ),
        "hold_claim_count": sum(
            row["최종실행상태"] == "HOLD" for row in claim_rows
        ),
        "final_reason_counts": dict(sorted(final_reasons.items())),
        "source_context_status": "FULL_ARTICLE_BODY_JOINED_NOT_REPARSED",
        "verdict_status": replay_summary.get("verdict_status", ""),
        "accuracy_status": replay_summary.get("accuracy_status", ""),
        "secret_handling": "NO_SECRET_READ_OR_RECORDED",
    }
    summary_path = output_dir / "summary.json"
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    workbook_path = output_dir / OUTPUT_NAMES[0]
    _write_workbook(
        workbook_path,
        article_rows=article_rows,
        claim_rows=claim_rows,
        query_rows=query_rows,
        reason_codes=sorted(final_reasons),
        inputs=(
            articles_csv,
            claim_ledger_csv,
            subtype_csv,
            claim_values_csv,
            metadata_snapshots_jsonl,
            value_snapshots_jsonl,
            replay_summary_json,
        ),
    )
    _write_manifest(
        output_dir,
        inputs=(
            articles_csv,
            claim_ledger_csv,
            subtype_csv,
            claim_values_csv,
            metadata_snapshots_jsonl,
            value_snapshots_jsonl,
            replay_summary_json,
        ),
    )
    return summary


def refresh_output_manifest(output_dir: Path) -> dict[str, Any]:
    """Refresh output hashes after an office suite recalculates the workbook."""
    manifest_path = output_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["outputs"] = {
        name: _file_record(output_dir / name) for name in OUTPUT_NAMES
    }
    manifest["manifest_refreshed_at"] = _now()
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest


def _build_article_rows(
    *,
    article_ids: set[str],
    article_by_id: Mapping[str, Mapping[str, str]],
    claims_by_article: Mapping[str, Sequence[Mapping[str, str]]],
    source_path: Path,
) -> list[dict[str, Any]]:
    output = []
    for article_id in sorted(article_ids):
        source = article_by_id[article_id]
        body = _text(source.get("본문"))
        claims = claims_by_article[article_id]
        exact_count = sum(_text(row.get("원문")) in body for row in claims)
        common = {
            "기사번호": article_id,
            "작성일": _text(source.get("작성일")),
            "제목": _text(source.get("제목")),
            "URL": _text(source.get("URL")),
            "검색구분레이블": _text(source.get("검색구분레이블")),
            "원본CSV행": _text(source.get("원본CSV행")),
            "파일명": _text(source.get("파일명")),
            "본문길이": len(body),
            "본문SHA256": hashlib.sha256(body.encode("utf-8")).hexdigest(),
            "Claim수": len(claims),
            "Claim원문정확일치수": exact_count,
            "기사문맥연결상태": (
                "COMPLETE" if body and exact_count == len(claims) else "PARTIAL"
            ),
            "원문데이터파일": str(source_path),
        }
        output.append({**common, "본문미리보기": body[:1200], "본문": body})
    return output


def _build_claim_rows(
    *,
    ledger_rows: Sequence[Mapping[str, str]],
    article_by_id: Mapping[str, Mapping[str, str]],
    subtype_by_claim: Mapping[str, Mapping[str, str]],
    value_by_claim: Mapping[str, Mapping[str, str]],
    value_snapshot_by_key: Mapping[tuple[Any, ...], Mapping[str, Any]],
) -> list[dict[str, Any]]:
    output = []
    for ledger in ledger_rows:
        claim_id = _text(ledger.get("Claim번호"))
        article_id = _text(ledger.get("기사번호"))
        article = article_by_id[article_id]
        body = _text(article.get("본문"))
        sentence = _text(ledger.get("원문"))
        position = body.find(sentence)
        subtype = subtype_by_claim[claim_id]
        official = value_by_claim[claim_id]
        snapshot = value_snapshot_by_key.get(_claim_value_key(official), {})
        official_status = _text(official.get("official_value_status"))
        reason = _text(official.get("reason_code"))
        final_status = (
            "OFFICIAL_VALUE_LINKED"
            if official_status == "OFFICIAL_VALUE_FETCHED"
            else "HOLD"
        )
        final_reason = (
            "OFFICIAL_VALUE_FETCHED"
            if final_status == "OFFICIAL_VALUE_LINKED"
            else reason or _text(ledger.get("최신결과사유")) or "UNSPECIFIED_HOLD"
        )
        row = {
            "기사번호": article_id,
            "Claim번호": claim_id,
            "문장번호": _text(ledger.get("문장번호")),
            "부모Claim번호": _text(ledger.get("부모Claim번호")),
            "작성일": _text(article.get("작성일")),
            "제목": _text(article.get("제목")),
            "URL": _text(article.get("URL")),
            "원문": sentence,
            "기사내원문시작위치": position,
            "기사원문정확일치": "YES" if position >= 0 else "NO",
            "앞문맥": body[max(0, position - 500):position] if position >= 0 else "",
            "뒤문맥": (
                body[position + len(sentence):position + len(sentence) + 500]
                if position >= 0
                else ""
            ),
            "분류탭": _text(subtype.get("탭")),
            "하위유형": _text(subtype.get("하위유형")),
            "분류속성": _text(subtype.get("속성")),
            "분류근거": _text(subtype.get("분류근거")),
            "12개항목상태": _text(ledger.get("12개항목상태")),
            "12개항목공식조회가능": _text(ledger.get("12개항목공식조회가능")),
            "통계표검색시도": _text(ledger.get("통계표검색시도")),
            "항목정보조회시도": _text(ledger.get("항목정보조회시도")),
            "기간정보조회시도": _text(ledger.get("기간정보조회시도")),
            "현재상태": _text(ledger.get("현재상태")),
            "현재중단단계": _text(ledger.get("현재중단단계")),
            "현재사유": _text(ledger.get("현재사유")),
            "대표문제": _text(ledger.get("대표문제")),
            "보조문제": _text(ledger.get("보조문제")),
            "다음실행단계": _text(ledger.get("다음실행단계")),
            "최신결과상태": _text(ledger.get("최신결과상태")),
            "최신결과단계": _text(ledger.get("최신결과단계")),
            "최신결과사유": _text(ledger.get("최신결과사유")),
            "최신개선판정": _text(ledger.get("최신개선판정")),
            "concept_id": _text(official.get("concept_id")),
            "standard_key": _text(official.get("standard_key")),
            "지표": _text(official.get("indicator")),
            "Claim단위": _text(official.get("claim_unit")),
            "Claim기간": _text(official.get("claim_time")),
            "Claim주기": _text(official.get("claim_frequency")),
            "계산유형": _text(official.get("calculation_type")),
            "registry_signature": _text(official.get("registry_signature")),
            "후보표수": _text(official.get("candidate_count")),
            "좌표상태": _text(official.get("coordinate_status")),
            "좌표사유": reason,
            "좌표출처": _text(official.get("coordinate_provenance")),
            "기관ID": _text(official.get("org_id")),
            "통계표ID": _text(official.get("table_id")),
            "항목ID": _text(official.get("item_id")),
            "분류코드": _text(official.get("object_codes")),
            "기간유형": _text(official.get("period_type")),
            "조회기간": _text(official.get("period")),
            "KOSIS값조회시도": "YES" if snapshot else "NO",
            "KOSIS조회시각": _text(snapshot.get("retrieved_at")),
            "KOSIS조회경로": _text(snapshot.get("retrieval_source")),
            "KOSIS조회상태": _text(snapshot.get("status")),
            "KOSIS오류코드": _text(snapshot.get("error_code")),
            "KOSIS응답SHA256": _text(snapshot.get("response_sha256")),
            "공식값상태": official_status,
            "공식값": _text(official.get("official_value")),
            "공식단위": _text(official.get("official_unit")),
            "Verdict상태": _text(official.get("verdict_status")),
            "최종실행상태": final_status,
            "최종성공실패사유": final_reason,
            "최종중단단계": _terminal_stage(final_status, final_reason),
        }
        output.append(row)
    return sorted(output, key=lambda row: row["Claim번호"])


def _build_query_rows(
    *,
    metadata_rows: Sequence[Mapping[str, Any]],
    value_snapshot_rows: Sequence[Mapping[str, Any]],
    claim_rows: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    claims_by_table: dict[str, list[str]] = defaultdict(list)
    claims_by_value_key: dict[tuple[Any, ...], list[str]] = defaultdict(list)
    for claim in claim_rows:
        table_id = _text(claim.get("통계표ID"))
        if table_id:
            claims_by_table[table_id].append(_text(claim.get("Claim번호")))
        key = _audit_claim_value_key(claim)
        if key[1]:
            claims_by_value_key[key].append(_text(claim.get("Claim번호")))

    output = []
    for snapshot in metadata_rows:
        request = snapshot.get("request", {})
        table_id = _text(request.get("table_id"))
        linked = sorted(set(claims_by_table.get(table_id, [])))
        response = snapshot.get("response", [])
        output.append(
            {
                "조회종류": "METADATA",
                "조회키": "|".join(
                    (
                        _text(request.get("org_id")),
                        table_id,
                        _text(request.get("meta_type")),
                    )
                ),
                "기관ID": _text(request.get("org_id")),
                "통계표ID": table_id,
                "메타유형": _text(request.get("meta_type")),
                "항목ID": "",
                "기간유형": "",
                "시작기간": "",
                "종료기간": "",
                "분류코드": "",
                "조회시각": _text(snapshot.get("retrieved_at")),
                "조회경로": _text(snapshot.get("retrieval_source")),
                "조회상태": _text(snapshot.get("status")),
                "오류코드": _text(snapshot.get("error_code")),
                "응답행수": len(response) if isinstance(response, list) else 0,
                "응답SHA256": _text(snapshot.get("response_sha256")),
                "연결Claim수": len(linked),
                "연결Claim번호": " | ".join(linked),
                "응답요약": _metadata_preview(response),
            }
        )
    for snapshot in value_snapshot_rows:
        request = snapshot.get("request", {})
        key = _snapshot_value_key(snapshot)
        linked = sorted(set(claims_by_value_key.get(key, [])))
        response = snapshot.get("response", [])
        output.append(
            {
                "조회종류": "VALUE",
                "조회키": _value_key_text(key),
                "기관ID": _text(request.get("org_id")),
                "통계표ID": _text(request.get("table_id")),
                "메타유형": "",
                "항목ID": _text(request.get("item_id")),
                "기간유형": _text(request.get("period_type")),
                "시작기간": _text(request.get("start_period")),
                "종료기간": _text(request.get("end_period")),
                "분류코드": " | ".join(_object_codes(request.get("object_codes"))),
                "조회시각": _text(snapshot.get("retrieved_at")),
                "조회경로": _text(snapshot.get("retrieval_source")),
                "조회상태": _text(snapshot.get("status")),
                "오류코드": _text(snapshot.get("error_code")),
                "응답행수": len(response) if isinstance(response, list) else 0,
                "응답SHA256": _text(snapshot.get("response_sha256")),
                "연결Claim수": len(linked),
                "연결Claim번호": " | ".join(linked),
                "응답요약": _value_preview(response),
            }
        )
    return sorted(output, key=lambda row: (row["조회종류"], row["조회키"]))


def _write_workbook(
    path: Path,
    *,
    article_rows: Sequence[Mapping[str, Any]],
    claim_rows: Sequence[Mapping[str, Any]],
    query_rows: Sequence[Mapping[str, Any]],
    reason_codes: Sequence[str],
    inputs: Sequence[Path],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "요약"
    article_sheet = workbook.create_sheet("전체 기사")
    claim_sheet = workbook.create_sheet("Claim 실행원장")
    query_sheet = workbook.create_sheet("KOSIS 조회기록")
    reason_sheet = workbook.create_sheet("사유 집계")
    info_sheet = workbook.create_sheet("실행정보")

    _append_table(article_sheet, ARTICLE_COLUMNS, article_rows)
    _append_table(claim_sheet, CLAIM_COLUMNS, claim_rows)
    _append_table(query_sheet, QUERY_COLUMNS, query_rows)

    article_status_col = _column_letter(ARTICLE_COLUMNS, "기사문맥연결상태")
    claim_context_col = _column_letter(CLAIM_COLUMNS, "기사원문정확일치")
    claim_final_col = _column_letter(CLAIM_COLUMNS, "최종실행상태")
    claim_reason_col = _column_letter(CLAIM_COLUMNS, "최종성공실패사유")
    query_kind_col = _column_letter(QUERY_COLUMNS, "조회종류")
    query_status_col = _column_letter(QUERY_COLUMNS, "조회상태")

    summary.append(["전체 기사·Claim·KOSIS 전수 실행 감사 원장"])
    summary.append(["지표", "값", "해석"])
    summary_rows = (
        (
            "전체 기사",
            f"=COUNTA('전체 기사'!A2:A{len(article_rows)+1})",
            "1,542 Claim이 연결된 고유 기사",
        ),
        (
            "기사 문맥 완전 연결",
            f'=COUNTIF(\'전체 기사\'!{article_status_col}2:{article_status_col}{len(article_rows)+1},"COMPLETE")',
            "기사 본문 안에서 해당 기사의 Claim 원문이 모두 확인됨",
        ),
        (
            "전체 Claim",
            f"=COUNTA('Claim 실행원장'!B2:B{len(claim_rows)+1})",
            "감사 원장에 기록된 Claim",
        ),
        (
            "기사 원문 정확 연결 Claim",
            f'=COUNTIF(\'Claim 실행원장\'!{claim_context_col}2:{claim_context_col}{len(claim_rows)+1},"YES")',
            "원문 Claim이 전체 기사 본문에 정확히 존재함",
        ),
        (
            "KOSIS metadata 조회 기록",
            f'=COUNTIF(\'KOSIS 조회기록\'!{query_kind_col}2:{query_kind_col}{len(query_rows)+1},"METADATA")',
            "ITM/PRD 구조정보 조회 또는 cache 재사용 기록",
        ),
        (
            "KOSIS 공식값 조회 기록",
            f'=COUNTIF(\'KOSIS 조회기록\'!{query_kind_col}2:{query_kind_col}{len(query_rows)+1},"VALUE")',
            "고유 공식 좌표 단위 조회 또는 cache 재사용 기록",
        ),
        (
            "KOSIS 조회 성공 기록",
            f'=COUNTIF(\'KOSIS 조회기록\'!{query_status_col}2:{query_status_col}{len(query_rows)+1},"SUCCESS")',
            "응답 형식 검증까지 성공한 조회 기록",
        ),
        (
            "공식값 연결 Claim",
            f'=COUNTIF(\'Claim 실행원장\'!{claim_final_col}2:{claim_final_col}{len(claim_rows)+1},"OFFICIAL_VALUE_LINKED")',
            "정확도가 아니라 공식값 연결 coverage",
        ),
        (
            "HOLD Claim",
            f'=COUNTIF(\'Claim 실행원장\'!{claim_final_col}2:{claim_final_col}{len(claim_rows)+1},"HOLD")',
            "중단 단계와 이유를 보존한 Claim",
        ),
    )
    for row in summary_rows:
        summary.append(row)
    summary.merge_cells("A1:C1")
    summary["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")

    reason_sheet.append(["최종성공실패사유", "Claim수"])
    for index, reason in enumerate(reason_codes, start=2):
        reason_sheet.cell(index, 1, reason)
        reason_sheet.cell(
            index,
            2,
            f'=COUNTIF(\'Claim 실행원장\'!{claim_reason_col}2:{claim_reason_col}{len(claim_rows)+1},A{index})',
        )

    info_sheet.append(["항목", "내용"])
    info_rows = [
        ("생성시각(UTC)", _now()),
        ("기사 범위", "1,542 Claim이 연결된 고유 기사 전체"),
        ("성공 정의", "OFFICIAL_VALUE_LINKED는 공식값 연결 성공이며 정확도 증명이 아님"),
        ("HOLD 정의", "중단 단계와 reason code가 기록된 안전한 미확정 상태"),
        ("기사문맥 상태", "전체 기사 본문은 연결했지만 이 산출물에서 Claim을 재추출하지 않음"),
        ("비밀정보", "API 키 값은 읽거나 기록하지 않음"),
    ]
    for source in inputs:
        info_rows.append((f"입력파일:{source.name}", json.dumps(_file_record(source), ensure_ascii=False)))
    for row in info_rows:
        info_sheet.append(row)

    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    _style_header_row(summary, 2)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 65
    info_sheet.column_dimensions["A"].width = 34
    info_sheet.column_dimensions["B"].width = 110
    for row in claim_sheet.iter_rows(min_row=2):
        for column_name in ("원문", "앞문맥", "뒤문맥", "12개항목상태", "분류근거"):
            row[CLAIM_COLUMNS.index(column_name)].alignment = Alignment(
                vertical="top", wrap_text=True
            )
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(path)


def _append_table(
    sheet: Any,
    columns: Sequence[str],
    rows: Sequence[Mapping[str, Any]],
) -> None:
    sheet.append(list(columns))
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _style_sheet(sheet: Any) -> None:
    _style_header_row(sheet, 1)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=9)
            if cell.alignment is None or not cell.alignment.wrap_text:
                cell.alignment = Alignment(vertical="top")
    sheet.sheet_view.showGridLines = False
    for index, cell in enumerate(sheet[1], start=1):
        width = 14
        name = _text(cell.value)
        if any(token in name for token in ("원문", "문맥", "사유", "근거", "번호")):
            width = 32
        if name in {"제목", "URL", "12개항목상태", "registry_signature", "응답요약"}:
            width = 48
        sheet.column_dimensions[get_column_letter(index)].width = width


def _style_header_row(sheet: Any, row_number: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[row_number]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _write_manifest(output_dir: Path, *, inputs: Sequence[Path]) -> None:
    manifest = {
        "schema_version": "clafact_full_article_execution_audit_v1",
        "created_at": _now(),
        "inputs": {str(path): _file_record(path) for path in inputs},
        "outputs": {name: _file_record(output_dir / name) for name in OUTPUT_NAMES},
        "secrets": {"api_key": "NOT_READ_NOT_RECORDED"},
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    output = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                output.append(json.loads(line))
    return output


def _write_csv(
    path: Path,
    rows: Sequence[Mapping[str, Any]],
    columns: Sequence[str],
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _require_unique(
    rows: Sequence[Mapping[str, Any]], column: str, label: str
) -> None:
    values = [_text(row.get(column)) for row in rows]
    if any(not value for value in values):
        raise ValueError(f"{label} has blank {column}")
    if len(values) != len(set(values)):
        raise ValueError(f"{label} has duplicate {column}")


def _claim_value_key(row: Mapping[str, Any]) -> tuple[Any, ...]:
    return (
        _text(row.get("org_id")),
        _text(row.get("table_id")),
        _text(row.get("item_id")),
        _text(row.get("period_type")),
        _text(row.get("period")),
        _text(row.get("period")),
        _object_codes(row.get("object_codes")),
    )


def _audit_claim_value_key(row: Mapping[str, Any]) -> tuple[Any, ...]:
    period = _text(row.get("조회기간"))
    return (
        _text(row.get("기관ID")),
        _text(row.get("통계표ID")),
        _text(row.get("항목ID")),
        _text(row.get("기간유형")),
        period,
        period,
        _object_codes(row.get("분류코드")),
    )


def _snapshot_value_key(row: Mapping[str, Any]) -> tuple[Any, ...]:
    request = row.get("request", {})
    return (
        _text(request.get("org_id")),
        _text(request.get("table_id")),
        _text(request.get("item_id")),
        _text(request.get("period_type")),
        _text(request.get("start_period")),
        _text(request.get("end_period")),
        _object_codes(request.get("object_codes")),
    )


def _object_codes(value: Any) -> tuple[str, ...]:
    if isinstance(value, (list, tuple)):
        return tuple(_text(item) for item in value if _text(item))
    text = _text(value)
    if not text:
        return ()
    return tuple(part.strip() for part in text.split("|") if part.strip())


def _value_key_text(key: tuple[Any, ...]) -> str:
    return "|".join(
        [*map(str, key[:6]), " + ".join(map(str, key[6]))]
    )


def _terminal_stage(final_status: str, reason: str) -> str:
    if final_status == "OFFICIAL_VALUE_LINKED":
        return "OFFICIAL_VALUE"
    if reason.startswith("KOSIS_VALUE"):
        return "KOSIS_VALUE"
    if reason.startswith("CONCEPT_"):
        return "CONCEPT"
    if reason in {"NO_REGISTRY_CANDIDATE", "MULTIPLE_OFFICIAL_COORDINATES_READY"}:
        return "REGISTRY"
    if reason.startswith("EVIDENCE_") or reason in {"TIME_NOT_AVAILABLE", "UNIT_CONFLICT"}:
        return "EVIDENCE_CELL"
    return "HOLD"


def _metadata_preview(response: Any) -> str:
    if not isinstance(response, list):
        return ""
    previews = []
    for row in response[:5]:
        if not isinstance(row, Mapping):
            continue
        if row.get("ITM_ID"):
            previews.append(
                f"{_text(row.get('OBJ_NM'))}:{_text(row.get('ITM_NM'))}"
                f"[{_text(row.get('ITM_ID'))}]"
            )
        elif row.get("STRT_PRD_DE"):
            previews.append(
                f"{_text(row.get('PRD_SE'))}:"
                f"{_text(row.get('STRT_PRD_DE'))}-{_text(row.get('END_PRD_DE'))}"
            )
    return " | ".join(previews)


def _value_preview(response: Any) -> str:
    if not isinstance(response, list):
        return ""
    previews = []
    for row in response[:3]:
        if isinstance(row, Mapping):
            previews.append(
                f"{_text(row.get('TBL_NM'))} / {_text(row.get('ITM_NM'))} / "
                f"{_text(row.get('PRD_DE'))} / {_text(row.get('DT'))} "
                f"{_text(row.get('UNIT_NM'))}"
            )
    return " | ".join(previews)


def _article_id(value: Any) -> str:
    digits = "".join(character for character in _text(value) if character.isdigit())
    return f"A{int(digits):05d}" if digits else ""


def _column_letter(columns: Sequence[str], name: str) -> str:
    return get_column_letter(columns.index(name) + 1)


def _file_record(path: Path) -> dict[str, Any]:
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--articles-csv", type=Path)
    parser.add_argument("--claim-ledger-csv", type=Path)
    parser.add_argument("--subtype-csv", type=Path)
    parser.add_argument("--claim-values-csv", type=Path)
    parser.add_argument("--metadata-snapshots-jsonl", type=Path)
    parser.add_argument("--value-snapshots-jsonl", type=Path)
    parser.add_argument("--replay-summary-json", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--expected-article-count", type=int, default=499)
    parser.add_argument("--expected-claim-count", type=int, default=1542)
    parser.add_argument("--refresh-manifest-only", action="store_true")
    args = parser.parse_args()
    if args.refresh_manifest_only:
        result = refresh_output_manifest(args.output_dir)
    else:
        required = {
            "articles_csv": args.articles_csv,
            "claim_ledger_csv": args.claim_ledger_csv,
            "subtype_csv": args.subtype_csv,
            "claim_values_csv": args.claim_values_csv,
            "metadata_snapshots_jsonl": args.metadata_snapshots_jsonl,
            "value_snapshots_jsonl": args.value_snapshots_jsonl,
            "replay_summary_json": args.replay_summary_json,
        }
        missing = sorted(name for name, value in required.items() if value is None)
        if missing:
            parser.error(f"missing required arguments: {', '.join(missing)}")
        result = build(
            **required,
            output_dir=args.output_dir,
            expected_article_count=args.expected_article_count,
            expected_claim_count=args.expected_claim_count,
        )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
