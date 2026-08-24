"""Run an offline structural Hard Guard replay without reading article text."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
from typing import Any, Sequence

from core.data_loader import load_kosis_catalog, load_period_availability_snapshot
from core.r3_hard_guard_readiness import evaluate_hard_guard_readiness
from schemas.candidate import KosisCandidateSchema
from schemas.period_availability import PeriodAvailabilitySnapshot


OUTPUT_COLUMNS = (
    "claim_id",
    "split",
    "guard_route_status",
    "attached_candidate_count",
    "normalized_candidate_count",
    "structural_pass_count",
    "surviving_candidate_tbl_ids",
    "reject_code_counts",
    "claim_schema_error",
    "source_context_status",
    "target_value_role_status",
    "selection_status",
    "next_system_gate",
)


def run(
    *,
    ledger_xlsx: Path,
    candidate_attachment_csv: Path,
    output_dir: Path,
    catalog_path: Path | None = None,
    period_snapshot_path: Path | None = None,
    period_manifest_path: Path | None = None,
    catalog_candidates: Sequence[KosisCandidateSchema] | None = None,
    period_availability: PeriodAvailabilitySnapshot | None = None,
    sheet_name: str = "1542건 원장",
) -> dict[str, Any]:
    """Write a count-safe Guard artifact to a new output directory."""
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(f"output directory is not empty: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)
    if catalog_candidates is None:
        if catalog_path is None:
            raise ValueError("catalog_path is required when catalog_candidates is absent")
        catalog_candidates = load_kosis_catalog(catalog_path)
    if period_availability is None and period_snapshot_path and period_manifest_path:
        period_availability = load_period_availability_snapshot(
            period_snapshot_path,
            period_manifest_path,
        )

    claims = read_claim_slots_without_article_text(ledger_xlsx, sheet_name=sheet_name)
    candidate_rows = _read_csv(candidate_attachment_csv)
    catalog = {candidate.tbl_id: candidate for candidate in catalog_candidates}
    rows, summary = evaluate_hard_guard_readiness(
        candidate_rows,
        claim_slots_by_id=claims,
        catalog_by_table_id=catalog,
        period_availability=period_availability,
    )

    result_csv = output_dir / "hard_guard_readiness.csv"
    summary_json = output_dir / "summary.json"
    _write_csv(result_csv, rows)
    _write_json(summary_json, summary)
    manifest = {
        "artifact": "r3_hard_guard_readiness_v1",
        "inputs": {
            "ledger_xlsx": _file_record(ledger_xlsx),
            "candidate_attachment_csv": _file_record(candidate_attachment_csv),
            "catalog_json": _optional_file_record(catalog_path),
            "period_snapshot": _optional_file_record(period_snapshot_path),
            "period_manifest": _optional_file_record(period_manifest_path),
        },
        "input_contract": {
            "article_text_column": "EXCLUDED_NOT_READ",
            "sheet_name": sheet_name,
            "slot_columns": [
                "indicator", "value", "unit", "time", "frequency", "region",
                "population", "dimension", "comparison", "calculation", "condition",
                "source_hint",
            ],
        },
        "outputs": {
            "hard_guard_readiness_csv": _file_record(result_csv),
            "summary_json": _file_record(summary_json),
        },
        "evaluation_boundary": summary["accuracy_status"],
    }
    _write_json(output_dir / "manifest.json", manifest)
    return summary


def read_claim_slots_without_article_text(
    path: Path, *, sheet_name: str
) -> dict[str, dict[str, object]]:
    """Read disjoint column blocks so the sentence column is never requested."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    left_header = next(
        worksheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=6, values_only=True)
    )
    slot_header = next(
        worksheet.iter_rows(min_row=1, max_row=1, min_col=8, max_col=19, values_only=True)
    )
    status_header = next(
        worksheet.iter_rows(min_row=1, max_row=1, min_col=20, max_col=26, values_only=True)
    )
    expected_left = ("claim_id", "article_id", "article_date", "split", "claim_type_original", "gold_time_source")
    expected_slots = (
        "indicator", "value", "unit", "time", "frequency", "region", "population",
        "dimension", "comparison", "calculation", "condition", "source_hint",
    )
    if tuple(_header(value) for value in left_header) != expected_left:
        raise ValueError("unexpected ledger identity columns")
    if tuple(_header(value) for value in slot_header) != expected_slots:
        raise ValueError("unexpected ledger slot columns")
    if _header(status_header[5]) != "target_value_role_status":
        raise ValueError("target_value_role_status column is missing")

    left_rows = worksheet.iter_rows(min_row=2, min_col=1, max_col=6, values_only=True)
    slot_rows = worksheet.iter_rows(min_row=2, min_col=8, max_col=19, values_only=True)
    status_rows = worksheet.iter_rows(min_row=2, min_col=20, max_col=26, values_only=True)
    output: dict[str, dict[str, object]] = {}
    for left, slots, statuses in zip(left_rows, slot_rows, status_rows, strict=True):
        claim_id = str(left[0] or "").strip()
        if not claim_id:
            continue
        if claim_id in output:
            raise ValueError(f"duplicate claim_id in ledger: {claim_id}")
        record = {
            name: _slot_value(name, value)
            for name, value in zip(expected_slots, slots, strict=True)
        }
        record.update(
            {
                "claim_id": claim_id,
                "split": str(left[3] or "").strip(),
                "target_value_role_status": str(statuses[5] or "").strip(),
            }
        )
        output[claim_id] = record
    workbook.close()
    return output


def _slot_value(name: str, value: object) -> object:
    if name in {"dimension", "comparison", "condition"}:
        if value in (None, ""):
            return None
        parsed = json.loads(str(value)) if isinstance(value, str) else value
        if not isinstance(parsed, dict):
            raise ValueError(f"{name} must decode to an object")
        return parsed or None
    return value


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def _optional_file_record(path: Path | None) -> dict[str, object] | None:
    return _file_record(path) if path is not None else None


def _file_record(path: Path) -> dict[str, object]:
    return {
        "path": str(path.resolve()),
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def _header(value: object) -> str:
    return str(value or "").lstrip("\ufeff").strip()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--ledger-xlsx", type=Path, required=True)
    parser.add_argument("--candidate-attachment-csv", type=Path, required=True)
    parser.add_argument("--catalog-json", type=Path, required=True)
    parser.add_argument("--period-snapshot", type=Path, required=True)
    parser.add_argument("--period-manifest", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--sheet-name", default="1542건 원장")
    args = parser.parse_args()
    summary = run(
        ledger_xlsx=args.ledger_xlsx,
        candidate_attachment_csv=args.candidate_attachment_csv,
        catalog_path=args.catalog_json,
        period_snapshot_path=args.period_snapshot,
        period_manifest_path=args.period_manifest,
        output_dir=args.output_dir,
        sheet_name=args.sheet_name,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
