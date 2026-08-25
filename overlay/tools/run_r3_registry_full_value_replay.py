"""Reinforce a provisional KOSIS registry and link official values to a full ledger."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import time
from collections import Counter, defaultdict
from collections.abc import Callable, Iterable, Mapping, Sequence
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from core.kosis_openapi_transport import get_meta
from core.kosis_value_transport import get_parameter_data
from core.r3_official_value_pilot import (
    ACCURACY_STATUS,
    VERDICT_STATUS,
    PilotTarget,
    evaluate_official_value_rows,
    prepare_official_coordinate,
)
from core.r3_registry_full_value_replay import (
    RegistryCandidate,
    build_composite_signature,
    choose_coordinate,
    coordinate_cache_key,
    merge_candidate_identities,
    prepare_registered_claim_coordinate,
    profile_matches_claim,
    rank_catalog_candidates,
)
from tools.run_r3_hard_guard_readiness import read_claim_slots_without_article_text


RESULT_COLUMNS = (
    "claim_id",
    "split",
    "concept_id",
    "standard_key",
    "indicator",
    "claim_unit",
    "claim_time",
    "claim_frequency",
    "calculation_type",
    "registry_signature",
    "candidate_count",
    "coordinate_status",
    "reason_code",
    "coordinate_provenance",
    "org_id",
    "table_id",
    "item_id",
    "object_codes",
    "period_type",
    "period",
    "official_value_status",
    "official_value",
    "official_unit",
    "value_response_sha256",
    "verdict_status",
)


def run(
    *,
    ledger_xlsx: Path,
    concept_xlsx: Path,
    candidate_attachment_csv: Path,
    candidate_identity_jsonl: Path,
    catalog_json: Path,
    registered_coordinates_json: Path,
    member_codes_json: Path,
    output_dir: Path,
    api_key: str,
    metadata_fetcher: Callable[..., Any] = get_meta,
    value_fetcher: Callable[..., Any] = get_parameter_data,
    metadata_cache_paths: Sequence[Path] = (),
    value_cache_paths: Sequence[Path] = (),
    requests_per_minute: float = 150,
    catalog_top_k: int = 5,
    ledger_sheet_name: str = "1542건 원장",
    concept_sheet_name: str = "03_1542_Concept_전체",
) -> dict[str, Any]:
    """Process every joined Claim while calling values only for exact cells."""
    if not api_key:
        raise RuntimeError("KOSIS_API_KEY_NOT_CONFIGURED")
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(f"output directory is not empty: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)

    claims = read_claim_slots_without_article_text(
        ledger_xlsx, sheet_name=ledger_sheet_name
    )
    concepts = _read_concepts_without_article_text(
        concept_xlsx, sheet_name=concept_sheet_name
    )
    if set(claims) != set(concepts):
        raise ValueError(
            "CLAIM_CONCEPT_JOIN_MISMATCH: "
            f"claims={len(claims)} concepts={len(concepts)} "
            f"joined={len(set(claims).intersection(concepts))}"
        )
    attachment_by_claim = {
        row["claim_id"]: row for row in _read_csv(candidate_attachment_csv)
    }
    identity_by_table = _read_identity_candidates(candidate_identity_jsonl)
    catalog_rows = _read_json_list(catalog_json)
    for row in catalog_rows:
        table_id = _text(row.get("TBL_ID") or row.get("tbl_id"))
        if table_id and table_id not in identity_by_table:
            identity_by_table[table_id] = {
                "org_id": _text(row.get("ORG_ID") or row.get("org_id")),
                "tbl_id": table_id,
                "tbl_name": _text(
                    row.get("TBL_NM_META")
                    or row.get("TBL_NM_INPUT")
                    or row.get("tbl_name")
                ),
            }
    registered_profiles = _read_json_list(registered_coordinates_json)
    member_codes = _read_json_list(member_codes_json)

    candidate_pools: dict[str, list[RegistryCandidate]] = {}
    for claim_id, slots in claims.items():
        attachment = attachment_by_claim.get(claim_id, {})
        official_search: list[RegistryCandidate] = []
        for table_id in _split_ids(attachment.get("candidate_tbl_ids")):
            identity = identity_by_table.get(table_id)
            if identity:
                official_search.append(
                    _identity_candidate(identity, source="OFFICIAL_SEARCH", score=0.8)
                )
        registered: list[RegistryCandidate] = []
        for profile in registered_profiles:
            if not profile_matches_claim(profile, slots):
                continue
            identity = identity_by_table.get(_text(profile.get("tbl_id")))
            if identity:
                registered.append(
                    _identity_candidate(
                        identity, source="REGISTERED_COORDINATE", score=1.0
                    )
                )
        local = rank_catalog_candidates(slots, catalog_rows, top_k=catalog_top_k)
        candidate_pools[claim_id] = merge_candidate_identities(
            registered, official_search, local
        )

    required_tables = {
        (candidate.org_id, candidate.table_id): candidate
        for candidates in candidate_pools.values()
        for candidate in candidates
    }
    metadata_cache = _load_metadata_cache(metadata_cache_paths)
    value_cache = _load_value_cache(value_cache_paths)
    limiter = _RateLimiter(requests_per_minute)
    metadata_live_calls = 0
    metadata_snapshots: list[dict[str, object]] = []
    metadata_by_table: dict[tuple[str, str], dict[str, object]] = {}
    metadata_path = output_dir / "metadata_snapshots.jsonl"
    with metadata_path.open("w", encoding="utf-8", newline="") as metadata_handle:
        for org_id, table_id in sorted(required_tables):
            table_meta: dict[str, object] = {}
            for meta_type in ("ITM", "PRD"):
                key = (org_id, table_id, meta_type)
                cached = metadata_cache.get(key)
                if cached is not None:
                    snapshot = _snapshot_from_cache(cached)
                else:
                    limiter.wait()
                    metadata_live_calls += 1
                    snapshot = _fetch_metadata_snapshot(
                        metadata_fetcher,
                        api_key=api_key,
                        org_id=org_id,
                        table_id=table_id,
                        meta_type=meta_type,
                    )
                metadata_snapshots.append(snapshot)
                metadata_handle.write(json.dumps(snapshot, ensure_ascii=False) + "\n")
                metadata_handle.flush()
                if snapshot["status"] == "SUCCESS":
                    table_meta[meta_type] = snapshot["response"]
            metadata_by_table[(org_id, table_id)] = table_meta

    profile_by_table = {
        _text(profile.get("tbl_id")): profile for profile in registered_profiles
    }
    decisions: dict[str, object] = {}
    coordinates_by_claim: dict[str, object] = {}
    results: list[dict[str, object]] = []
    result_index_by_claim: dict[str, int] = {}
    slots_by_claim = claims
    for claim_id in sorted(claims):
        slots = claims[claim_id]
        concept = concepts[claim_id]
        signature = build_composite_signature(slots, concept)
        candidates = candidate_pools[claim_id]
        if _text(concept.get("concept_status")) != "ASSIGNED":
            decision = None
            reason = "CONCEPT_REGISTRY_NOT_READY"
        elif not candidates:
            decision = None
            reason = "NO_REGISTRY_CANDIDATE"
        else:
            prepared = []
            for candidate in candidates:
                metadata = metadata_by_table.get(
                    (candidate.org_id, candidate.table_id), {}
                )
                item_rows = metadata.get("ITM")
                period_rows = metadata.get("PRD")
                target = PilotTarget(
                    claim_id=claim_id,
                    split=_text(slots.get("split")),
                    indicator=_text(slots.get("indicator")),
                    table_id=candidate.table_id,
                    table_name=candidate.table_name,
                    org_id=candidate.org_id,
                    catalog_scope=candidate.source,
                    candidate_claim_count=1,
                )
                if not isinstance(item_rows, list) or not isinstance(period_rows, list):
                    continue
                profile = profile_by_table.get(candidate.table_id)
                if candidate.source == "REGISTERED_COORDINATE" and profile:
                    coordinate = prepare_registered_claim_coordinate(
                        target,
                        slots,
                        profile,
                        member_codes,
                        item_rows,
                        period_rows,
                    )
                else:
                    coordinate = prepare_official_coordinate(
                        target, slots, item_rows, period_rows
                    )
                prepared.append(coordinate)
            decision = choose_coordinate(prepared)
            decisions[claim_id] = decision
            reason = decision.reason_code
            if decision.coordinate is not None:
                coordinates_by_claim[claim_id] = decision.coordinate

        coordinate = coordinates_by_claim.get(claim_id)
        result = {
            "claim_id": claim_id,
            "split": _text(slots.get("split")),
            "concept_id": _text(concept.get("concept_id")),
            "standard_key": _text(concept.get("standard_key")),
            "indicator": _text(slots.get("indicator")),
            "claim_unit": _text(slots.get("unit")),
            "claim_time": _text(slots.get("time")),
            "claim_frequency": _text(slots.get("frequency")),
            "calculation_type": _structured_type(slots.get("calculation")),
            "registry_signature": signature,
            "candidate_count": len(candidates),
            "coordinate_status": (
                decision.status if decision is not None else "HOLD_COORDINATE_UNRESOLVED"
            ),
            "reason_code": reason,
            "coordinate_provenance": (
                coordinate.target.catalog_scope if coordinate is not None else ""
            ),
            "org_id": coordinate.target.org_id if coordinate is not None else "",
            "table_id": coordinate.target.table_id if coordinate is not None else "",
            "item_id": coordinate.item_id if coordinate is not None else "",
            "object_codes": (
                " | ".join(coordinate.object_codes) if coordinate is not None else ""
            ),
            "period_type": coordinate.period_type if coordinate is not None else "",
            "period": coordinate.period if coordinate is not None else "",
            "official_value_status": "PENDING" if coordinate is not None else "HOLD",
            "official_value": "",
            "official_unit": "",
            "value_response_sha256": "",
            "verdict_status": VERDICT_STATUS,
        }
        result_index_by_claim[claim_id] = len(results)
        results.append(result)

    claims_by_cell: dict[tuple[object, ...], list[str]] = defaultdict(list)
    for claim_id, coordinate in coordinates_by_claim.items():
        claims_by_cell[coordinate_cache_key(coordinate)].append(claim_id)

    value_live_calls = 0
    value_cache_hits = 0
    value_snapshots: list[dict[str, object]] = []
    value_path = output_dir / "value_snapshots.jsonl"
    with value_path.open("w", encoding="utf-8", newline="") as value_handle:
        for key in sorted(claims_by_cell, key=lambda item: tuple(map(str, item))):
            claim_ids = claims_by_cell[key]
            representative_id = claim_ids[0]
            coordinate = coordinates_by_claim[representative_id]
            cached_value = value_cache.get(key)
            if cached_value is not None:
                value_cache_hits += 1
                snapshot = _snapshot_from_cache(cached_value)
            else:
                limiter.wait()
                value_live_calls += 1
                snapshot = _fetch_value_snapshot(
                    value_fetcher, api_key=api_key, coordinate=coordinate
                )
            value_snapshots.append(snapshot)
            value_handle.write(json.dumps(snapshot, ensure_ascii=False) + "\n")
            value_handle.flush()
            for claim_id in claim_ids:
                result = results[result_index_by_claim[claim_id]]
                if snapshot["status"] != "SUCCESS":
                    result["official_value_status"] = "HOLD"
                    result["reason_code"] = snapshot["error_code"]
                    continue
                evaluated = evaluate_official_value_rows(
                    coordinates_by_claim[claim_id],
                    slots_by_claim[claim_id],
                    snapshot["response"],
                )
                result.update(
                    {
                        "official_value_status": evaluated["official_value_status"],
                        "reason_code": evaluated["reason_code"],
                        "official_value": evaluated["official_value"],
                        "official_unit": evaluated["official_unit"],
                        "value_response_sha256": snapshot["response_sha256"],
                    }
                )

    registry_rows = _build_registry_rows(results)
    hold_counts = Counter(
        _text(row.get("reason_code")) or "UNSPECIFIED_HOLD"
        for row in results
        if row["official_value_status"] != "OFFICIAL_VALUE_FETCHED"
    )
    coordinate_status_counts = Counter(_text(row["coordinate_status"]) for row in results)
    value_status_counts = Counter(_text(row["official_value_status"]) for row in results)
    summary = {
        "artifact": "r3_registry_full_value_replay_v1",
        "input_claim_count": len(claims),
        "concept_join_count": len(concepts),
        "registry_signature_count": len(registry_rows),
        "candidate_table_count": len(required_tables),
        "metadata_snapshot_count": len(metadata_snapshots),
        "metadata_live_api_call_count": metadata_live_calls,
        "metadata_cache_hit_count": len(metadata_snapshots) - metadata_live_calls,
        "coordinate_ready_claim_count": len(coordinates_by_claim),
        "unique_official_cell_count": len(claims_by_cell),
        "unique_value_api_call_count": value_live_calls,
        "value_cache_hit_count": value_cache_hits,
        "official_value_linked_claim_count": sum(
            row["official_value_status"] == "OFFICIAL_VALUE_FETCHED" for row in results
        ),
        "coordinate_status_counts": dict(sorted(coordinate_status_counts.items())),
        "official_value_status_counts": dict(sorted(value_status_counts.items())),
        "hold_reason_counts": dict(sorted(hold_counts.items())),
        "source_context_status": "NOT_READ_ARTICLE_TEXT",
        "target_value_role_status": "NOT_USED_FOR_VERDICT",
        "official_value_scope": "CURRENT_RELEASE_FOR_CLAIM_PERIOD_ARTICLE_ASOF_UNCHECKED",
        "table_selection_status": "REGISTERED_OR_PROVISIONAL_UNIQUE_METADATA_ONLY",
        "verdict_status": VERDICT_STATUS,
        "accuracy_status": ACCURACY_STATUS,
    }
    _write_csv(output_dir / "claim_value_results.csv", results, RESULT_COLUMNS)
    _write_csv(
        output_dir / "registry.csv",
        registry_rows,
        tuple(registry_rows[0]) if registry_rows else ("registry_signature",),
    )
    _write_json(output_dir / "summary.json", summary)
    _write_workbook(
        output_dir / "CLAFACT_1542_공식값_조회결과.xlsx",
        summary=summary,
        registry_rows=registry_rows,
        result_rows=results,
        hold_counts=hold_counts,
    )
    manifest = {
        "artifact": "r3_registry_full_value_replay_manifest_v1",
        "generated_at": _now(),
        "inputs": {
            "ledger_xlsx": _file_record(ledger_xlsx),
            "concept_xlsx": _file_record(concept_xlsx),
            "candidate_attachment_csv": _file_record(candidate_attachment_csv),
            "candidate_identity_jsonl": _file_record(candidate_identity_jsonl),
            "catalog_json": _file_record(catalog_json),
            "registered_coordinates_json": _file_record(registered_coordinates_json),
            "member_codes_json": _file_record(member_codes_json),
            "metadata_cache_paths": [_file_record(path) for path in metadata_cache_paths],
            "value_cache_paths": [_file_record(path) for path in value_cache_paths],
        },
        "privacy": {
            "article_text": "EXCLUDED_NOT_READ",
            "article_url": "EXCLUDED_NOT_READ",
        },
        "secrets": {"kosis_api_key": "PRESENT_NOT_RECORDED"},
        "outputs": {
            name: _file_record(output_dir / name)
            for name in (
                "claim_value_results.csv",
                "registry.csv",
                "metadata_snapshots.jsonl",
                "value_snapshots.jsonl",
                "summary.json",
                "CLAFACT_1542_공식값_조회결과.xlsx",
            )
        },
    }
    _write_json(output_dir / "manifest.json", manifest)
    return summary


class _RateLimiter:
    def __init__(self, requests_per_minute: float) -> None:
        self.interval = 60 / requests_per_minute if requests_per_minute > 0 else 0
        self.last_call = 0.0

    def wait(self) -> None:
        if not self.interval:
            return
        now = time.monotonic()
        remaining = self.interval - (now - self.last_call)
        if remaining > 0:
            time.sleep(remaining)
        self.last_call = time.monotonic()


def _read_concepts_without_article_text(
    path: Path, *, sheet_name: str
) -> dict[str, dict[str, object]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    left_header = next(
        worksheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=2, values_only=True)
    )
    right_header = next(
        worksheet.iter_rows(min_row=1, max_row=1, min_col=4, max_col=23, values_only=True)
    )
    if tuple(_header(value) for value in left_header) != ("article_id", "sentence_id"):
        raise ValueError("unexpected concept identity columns")
    expected = (
        "indicator", "value", "unit", "time", "frequency", "region", "population",
        "dimension", "comparison", "calculation", "condition", "source_hint",
        "parse_status", "concept_id", "canonical_name", "standard_key",
        "concept_action", "concept_status", "concept_reason", "kosis_fit_status",
    )
    if tuple(_header(value) for value in right_header) != expected:
        raise ValueError("unexpected concept columns")
    left_rows = worksheet.iter_rows(min_row=2, min_col=1, max_col=2, values_only=True)
    right_rows = worksheet.iter_rows(min_row=2, min_col=4, max_col=23, values_only=True)
    output: dict[str, dict[str, object]] = {}
    for left, right in zip(left_rows, right_rows, strict=True):
        claim_id = f"{_text(left[0])}_{_text(left[1])}"
        if claim_id in output:
            raise ValueError(f"duplicate concept claim_id: {claim_id}")
        output[claim_id] = dict(zip(expected, right, strict=True))
    workbook.close()
    return output


def _read_identity_candidates(path: Path) -> dict[str, dict[str, object]]:
    output: dict[str, dict[str, object]] = {}
    with path.open("r", encoding="utf-8-sig") as handle:
        for line in handle:
            if not line.strip():
                continue
            payload = json.loads(line)
            candidates = payload.get("candidates", []) if isinstance(payload, Mapping) else []
            if isinstance(candidates, list):
                for candidate in candidates:
                    if not isinstance(candidate, Mapping):
                        continue
                    table_id = _text(candidate.get("tbl_id") or candidate.get("TBL_ID"))
                    if table_id:
                        output.setdefault(table_id, dict(candidate))
    return output


def _registered_aliases(
    profiles: Iterable[Mapping[str, object]],
) -> dict[str, list[Mapping[str, object]]]:
    output: dict[str, list[Mapping[str, object]]] = defaultdict(list)
    for profile in profiles:
        aliases = profile.get("indicator_aliases")
        if isinstance(aliases, list):
            for alias in aliases:
                output[_normalize(_text(alias))].append(profile)
    return output


def _structured_type(value: object) -> str:
    if isinstance(value, Mapping):
        return _text(value.get("type"))
    text = _text(value)
    if text.startswith("{"):
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError:
            return text
        if isinstance(parsed, Mapping):
            return _text(parsed.get("type"))
    return text


def _load_metadata_cache(
    paths: Sequence[Path],
) -> dict[tuple[str, str, str], dict[str, object]]:
    output: dict[tuple[str, str, str], dict[str, object]] = {}
    for path in paths:
        with path.open("r", encoding="utf-8-sig") as handle:
            for line in handle:
                if not line.strip():
                    continue
                row = json.loads(line)
                request = row.get("request", {})
                key = (
                    _text(request.get("org_id")),
                    _text(request.get("table_id")),
                    _text(request.get("meta_type")),
                )
                if all(key) and row.get("status") == "SUCCESS":
                    output.setdefault(key, row)
    return output


def _load_value_cache(
    paths: Sequence[Path],
) -> dict[tuple[object, ...], dict[str, object]]:
    output: dict[tuple[object, ...], dict[str, object]] = {}
    for path in paths:
        with path.open("r", encoding="utf-8-sig") as handle:
            for line in handle:
                if not line.strip():
                    continue
                row = json.loads(line)
                request = row.get("request", {})
                object_codes = request.get("object_codes", [])
                if not isinstance(object_codes, list):
                    continue
                key = (
                    _text(request.get("org_id")),
                    _text(request.get("table_id")),
                    _text(request.get("item_id")),
                    _text(request.get("period_type")),
                    _text(request.get("start_period")),
                    *(_text(code) for code in object_codes),
                )
                if all(key) and row.get("status") == "SUCCESS":
                    output.setdefault(key, row)
    return output


def _snapshot_from_cache(row: Mapping[str, object]) -> dict[str, object]:
    snapshot = dict(row)
    snapshot["retrieval_source"] = "LOCAL_CACHE_REUSED"
    return snapshot


def _fetch_metadata_snapshot(
    fetcher: Callable[..., Any],
    *,
    api_key: str,
    org_id: str,
    table_id: str,
    meta_type: str,
) -> dict[str, object]:
    request = {"org_id": org_id, "table_id": table_id, "meta_type": meta_type}
    try:
        response = fetcher(
            api_key, org_id, table_id, meta_type=meta_type
        )
        return {
            "request": request,
            "retrieved_at": _now(),
            "retrieval_source": "LIVE_KOSIS",
            "status": "SUCCESS",
            "error_code": "",
            "response_sha256": _payload_hash(response),
            "response": response,
        }
    except Exception as error:
        return {
            "request": request,
            "retrieved_at": _now(),
            "retrieval_source": "LIVE_KOSIS",
            "status": "HOLD",
            "error_code": _safe_error(error),
            "response_sha256": "",
            "response": [],
        }


def _fetch_value_snapshot(
    fetcher: Callable[..., Any],
    *,
    api_key: str,
    coordinate: Any,
) -> dict[str, object]:
    request = {
        "org_id": coordinate.target.org_id,
        "table_id": coordinate.target.table_id,
        "item_id": coordinate.item_id,
        "period_type": coordinate.period_type,
        "start_period": coordinate.period,
        "end_period": coordinate.period,
        "object_codes": list(coordinate.object_codes),
    }
    try:
        response = fetcher(
            api_key,
            coordinate.target.org_id,
            coordinate.target.table_id,
            coordinate.item_id,
            coordinate.period_type,
            coordinate.period,
            coordinate.period,
            list(coordinate.object_codes),
        )
        return {
            "request": request,
            "retrieved_at": _now(),
            "retrieval_source": "LIVE_KOSIS",
            "status": "SUCCESS",
            "error_code": "",
            "response_sha256": _payload_hash(response),
            "response": response,
        }
    except Exception as error:
        return {
            "request": request,
            "retrieved_at": _now(),
            "retrieval_source": "LIVE_KOSIS",
            "status": "HOLD",
            "error_code": _safe_error(error),
            "response_sha256": "",
            "response": [],
        }


def _build_registry_rows(results: Iterable[Mapping[str, object]]) -> list[dict[str, object]]:
    grouped: dict[str, list[Mapping[str, object]]] = defaultdict(list)
    for row in results:
        grouped[_text(row.get("registry_signature"))].append(row)
    output = []
    for signature, rows in sorted(grouped.items()):
        ready = [row for row in rows if _text(row.get("table_id"))]
        tables = sorted({_text(row.get("table_id")) for row in ready})
        items = sorted({_text(row.get("item_id")) for row in ready})
        provenance = sorted({_text(row.get("coordinate_provenance")) for row in ready})
        reasons = Counter(
            _text(row.get("reason_code"))
            for row in rows
            if _text(row.get("reason_code"))
        )
        if ready and provenance == ["REGISTERED_COORDINATE"]:
            status = "REGISTERED_COORDINATE_VALIDATED"
        elif len(tables) == 1 and ready:
            status = "PROVISIONAL_UNIQUE_METADATA_REGISTRY"
        elif ready:
            status = "PROVISIONAL_MULTI_TABLE_BY_CLAIM"
        else:
            status = "HOLD_REGISTRY_UNRESOLVED"
        output.append(
            {
                "registry_signature": signature,
                "claim_count": len(rows),
                "coordinate_ready_claim_count": len(ready),
                "official_value_linked_claim_count": sum(
                    row.get("official_value_status") == "OFFICIAL_VALUE_FETCHED"
                    for row in rows
                ),
                "registry_status": status,
                "coordinate_provenance": " | ".join(provenance),
                "table_ids": " | ".join(tables),
                "item_ids": " | ".join(items),
                "hold_reason_counts": json.dumps(
                    dict(sorted(reasons.items())), ensure_ascii=False
                ),
            }
        )
    return output


def _write_workbook(
    path: Path,
    *,
    summary: Mapping[str, object],
    registry_rows: list[Mapping[str, object]],
    result_rows: list[Mapping[str, object]],
    hold_counts: Mapping[str, int],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "요약"
    summary_sheet.append(["CLAFACT 1,542건 KOSIS 공식값 조회 결과", ""])
    summary_sheet.append(["주의", "공식값 조회 coverage이며 표·좌표·Verdict 정확도가 아님"])
    summary_sheet.append(["기사시점", "현재 공개본에서 Claim 기간을 조회; 기사 당시 수정 전 값은 미확인"])
    summary_sheet.append(["KOSIS 개발가이드", "https://kosis.kr/openapi/devGuide/devGuide_0201List.do"])
    summary_sheet.append([])
    for key, value in summary.items():
        if isinstance(value, Mapping):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        summary_sheet.append([key, value])

    registry_sheet = workbook.create_sheet("Registry")
    _append_mapping_rows(registry_sheet, registry_rows)
    results_sheet = workbook.create_sheet("Claim 공식값")
    _append_mapping_rows(results_sheet, result_rows, columns=RESULT_COLUMNS)
    holds_sheet = workbook.create_sheet("HOLD 사유")
    holds_sheet.append(["reason_code", "claim_count"])
    for reason, count in sorted(hold_counts.items(), key=lambda item: (-item[1], item[0])):
        holds_sheet.append([reason, count])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    sub_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if sheet.max_row and sheet.max_column:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = header_fill
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            maximum = max(len(str(cell.value or "")) for cell in column_cells[:200])
            sheet.column_dimensions[letter].width = min(max(maximum + 2, 12), 36)
    summary_sheet["A1"].font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    for row in range(2, 5):
        summary_sheet.cell(row=row, column=1).fill = sub_fill
        summary_sheet.cell(row=row, column=1).font = Font(name="Arial", size=10, bold=True)
    summary_sheet.freeze_panes = "A6"
    workbook.save(path)


def _append_mapping_rows(
    sheet: Any,
    rows: Sequence[Mapping[str, object]],
    *,
    columns: Sequence[str] | None = None,
) -> None:
    selected = tuple(columns or (tuple(rows[0]) if rows else ("status",)))
    sheet.append(list(selected))
    for row in rows:
        sheet.append([row.get(column, "") for column in selected])


def _identity_candidate(
    identity: Mapping[str, object], *, source: str, score: float
) -> RegistryCandidate:
    return RegistryCandidate(
        org_id=_text(identity.get("org_id") or identity.get("ORG_ID")),
        table_id=_text(identity.get("tbl_id") or identity.get("TBL_ID")),
        table_name=_text(
            identity.get("tbl_name")
            or identity.get("TBL_NM_META")
            or identity.get("TBL_NM_INPUT")
        ),
        source=source,
        score=score,
    )


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _read_json_list(path: Path) -> list[dict[str, object]]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(payload, list) or not all(isinstance(row, dict) for row in payload):
        raise ValueError(f"expected JSON list: {path}")
    return payload


def _write_csv(
    path: Path, rows: Sequence[Mapping[str, object]], columns: Sequence[str]
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _write_json(path: Path, payload: Mapping[str, object]) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def _payload_hash(payload: object) -> str:
    encoded = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _file_record(path: Path) -> dict[str, object]:
    return {
        "path": str(path.resolve()),
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def _safe_error(error: Exception) -> str:
    text = str(error).strip()
    if text.startswith("KOSIS_") or text.startswith("OFFICIAL_"):
        return text
    return type(error).__name__.upper()


def _split_ids(value: object) -> list[str]:
    return [part.strip() for part in _text(value).split("|") if part.strip()]


def _normalize(value: str) -> str:
    return "".join(character for character in value.casefold() if character.isalnum())


def _header(value: object) -> str:
    return _text(value).lstrip("\ufeff")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def main() -> int:
    from config.settings import Settings

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ledger-xlsx", type=Path, required=True)
    parser.add_argument("--concept-xlsx", type=Path, required=True)
    parser.add_argument("--candidate-attachment-csv", type=Path, required=True)
    parser.add_argument("--candidate-identity-jsonl", type=Path, required=True)
    parser.add_argument("--catalog-json", type=Path, required=True)
    parser.add_argument("--registered-coordinates-json", type=Path, required=True)
    parser.add_argument("--member-codes-json", type=Path, required=True)
    parser.add_argument("--metadata-cache-jsonl", type=Path, action="append", default=[])
    parser.add_argument("--value-cache-jsonl", type=Path, action="append", default=[])
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--requests-per-minute", type=float, default=150)
    parser.add_argument("--catalog-top-k", type=int, default=5)
    parser.add_argument("--allow-live-kosis", action="store_true")
    args = parser.parse_args()
    if not args.allow_live_kosis:
        raise RuntimeError("LIVE_KOSIS_APPROVAL_FLAG_REQUIRED")
    settings = Settings()
    summary = run(
        ledger_xlsx=args.ledger_xlsx,
        concept_xlsx=args.concept_xlsx,
        candidate_attachment_csv=args.candidate_attachment_csv,
        candidate_identity_jsonl=args.candidate_identity_jsonl,
        catalog_json=args.catalog_json,
        registered_coordinates_json=args.registered_coordinates_json,
        member_codes_json=args.member_codes_json,
        metadata_cache_paths=args.metadata_cache_jsonl,
        value_cache_paths=args.value_cache_jsonl,
        output_dir=args.output_dir,
        api_key=settings.kosis_api_key or "",
        requests_per_minute=args.requests_per_minute,
        catalog_top_k=args.catalog_top_k,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
